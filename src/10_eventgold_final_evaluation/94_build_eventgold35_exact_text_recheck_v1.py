from pathlib import Path
import hashlib

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path.cwd()

FINAL = (
    ROOT
    / "01_event_aware_v3"
    / "10_eventgold_final_evaluation"
)

EXT = FINAL / "eventgold35_supervised_extension_v1"

A_PATH = EXT / "EventGold35_ANNOTATOR_A_RECEIVED_ORIGINAL_v1.xlsx"
B_PATH = EXT / "EventGold35_ANNOTATOR_B_RECEIVED_ORIGINAL_v1.xlsx"

BLIND_PATH = EXT / "EventGold35_BLIND_MASTER_v1.csv"

AUDIT_PATH = EXT / "EventGold35_ANNOTATION_TEXT_MISMATCH_AUDIT_v1.csv"

OUT_A = EXT / "EventGold35_ANNOTATOR_A_RECHECK_8_v1.xlsx"
OUT_B = EXT / "EventGold35_ANNOTATOR_B_RECHECK_8_v1.xlsx"


FRAME = [
    "clinical innovation and diagnostic benefit",
    "health safety risk and medical harm",
    "misinformation, deepfakes, and fraud",
    "regulation, governance, and policy",
    "public trust and accountability",
    "access, inequality, and digital divide",
    "labor automation and professional displacement",
    "mental health and therapy risk",
    "other / unclear",
]

STANCE = [
    "optimistic about AI in health",
    "cautious or mixed about AI in health",
    "critical of AI risks in health",
    "focused on regulation and governance",
    "unclear / not applicable",
]

MISINFO = [
    "direct misinformation focus",
    "indirect misinformation concern",
    "no clear misinformation relation",
    "unclear / not applicable",
]


def clean(x):
    if pd.isna(x):
        return ""

    x = str(x).replace("\r", " ").replace("\n", " ")
    return " ".join(x.split()).strip()


def read_original(path):

    df = pd.read_excel(
        path,
        sheet_name="Annotation",
        header=3,
        dtype=str,
        keep_default_na=False,
    )

    required = [
        "annotation_id",
        "Title",
        "Excerpt",
    ]

    missing = [
        c for c in required
        if c not in df.columns
    ]

    if missing:
        raise RuntimeError(
            f"Missing columns in {path.name}: {missing}"
        )

    df = df[required].copy()

    for c in required:
        df[c] = df[c].map(clean)

    if len(df) != 35:
        raise RuntimeError(
            f"{path.name}: expected 35 rows."
        )

    return df


blind = pd.read_csv(
    BLIND_PATH,
    dtype=str,
    keep_default_na=False,
)

for c in blind.columns:
    blind[c] = blind[c].map(clean)

a = read_original(A_PATH)
b = read_original(B_PATH)


# ------------------------------------------------------------
# IDs and titles must still match the frozen packet.
# ------------------------------------------------------------

for name, df in [("A", a), ("B", b)]:

    if (
        df["annotation_id"].tolist()
        != blind["annotation_id"].tolist()
    ):
        raise RuntimeError(
            f"Reviewer {name}: annotation IDs/order mismatch."
        )

    if (
        df["Title"].tolist()
        != blind["title"].tolist()
    ):
        raise RuntimeError(
            f"Reviewer {name}: title mismatch."
        )


# ------------------------------------------------------------
# Detect excerpt deviations mechanically.
# No labels are consulted here.
# ------------------------------------------------------------

audit_rows = []

mismatch_a = []
mismatch_b = []

for i in range(len(blind)):

    annotation_id = blind.loc[i, "annotation_id"]

    frozen = blind.loc[i, "annotation_excerpt"]
    ea = a.loc[i, "Excerpt"]
    eb = b.loc[i, "Excerpt"]

    a_match = ea == frozen
    b_match = eb == frozen

    if not a_match:
        mismatch_a.append(annotation_id)

    if not b_match:
        mismatch_b.append(annotation_id)

    if (not a_match) or (not b_match):

        audit_rows.append({
            "annotation_id": annotation_id,
            "title": blind.loc[i, "title"],
            "annotator_A_excerpt_matches_frozen": a_match,
            "annotator_B_excerpt_matches_frozen": b_match,
            "frozen_excerpt_chars": len(frozen),
            "annotator_A_excerpt_chars": len(ea),
            "annotator_B_excerpt_chars": len(eb),
        })


if mismatch_a != mismatch_b:
    raise RuntimeError(
        "A and B do not have the same excerpt-mismatch IDs."
    )

if not mismatch_a:
    raise RuntimeError(
        "No excerpt mismatch detected; unexpected state."
    )


expected_ids = [
    "EG35_ANN_003",
    "EG35_ANN_007",
    "EG35_ANN_009",
    "EG35_ANN_011",
    "EG35_ANN_013",
    "EG35_ANN_016",
    "EG35_ANN_017",
    "EG35_ANN_021",
]

if mismatch_a != expected_ids:
    raise RuntimeError(
        "Mismatch set differs from independently audited "
        f"expected set.\nObserved: {mismatch_a}"
    )


audit = pd.DataFrame(audit_rows)

audit.to_csv(
    AUDIT_PATH,
    index=False,
    lineterminator="\n",
)


# ------------------------------------------------------------
# Build blinded 8-document corrective recheck workbooks.
# Previous labels are intentionally NOT included.
# ------------------------------------------------------------

recheck = blind[
    blind["annotation_id"].isin(expected_ids)
].copy()

# retain original frozen annotation order
recheck = recheck.set_index(
    "annotation_id"
).loc[expected_ids].reset_index()


def build_workbook(path, reviewer):

    wb = Workbook()

    ws = wb.active
    ws.title = "Annotation"

    ws["A1"] = (
        f"EventGold35 — Exact Frozen Text Recheck — Annotator {reviewer}"
    )
    ws["A1"].font = Font(
        bold=True,
        size=14,
    )

    ws.merge_cells("A1:H1")

    ws["A2"] = (
        "Re-annotate ONLY these 8 documents independently. "
        "These rows correct a presentation-text mismatch detected "
        "before agreement analysis. Do not consult your previous "
        "answers or the other annotator."
    )
    ws.merge_cells("A2:H2")
    ws["A2"].alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )

    headers = [
        "No.",
        "annotation_id",
        "Title",
        "Exact frozen excerpt",
        "Primary frame",
        "Stance toward AI in health",
        "Relation to misinformation",
        "Optional notes",
    ]

    for c, value in enumerate(
        headers,
        start=1,
    ):
        cell = ws.cell(
            row=4,
            column=c,
            value=value,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            "solid",
            fgColor="16324F",
        )

        cell.alignment = Alignment(
            wrap_text=True,
            vertical="center",
        )

    for j, row in recheck.iterrows():

        r = j + 5

        ws.cell(r, 1, j + 1)
        ws.cell(r, 2, row["annotation_id"])
        ws.cell(r, 3, row["title"])
        ws.cell(r, 4, row["annotation_excerpt"])

        # E:H deliberately blank.

        for c in range(1, 9):
            ws.cell(r, c).alignment = Alignment(
                wrap_text=True,
                vertical="top",
            )

    # Options sheet
    opt = wb.create_sheet("Options")

    opt["A1"] = "Primary frame"
    opt["B1"] = "Stance"
    opt["C1"] = "Misinformation"

    for i, x in enumerate(FRAME, start=2):
        opt.cell(i, 1, x)

    for i, x in enumerate(STANCE, start=2):
        opt.cell(i, 2, x)

    for i, x in enumerate(MISINFO, start=2):
        opt.cell(i, 3, x)

    # Dropdowns
    dv_frame = DataValidation(
        type="list",
        formula1="'Options'!$A$2:$A$10",
        allow_blank=False,
    )

    dv_stance = DataValidation(
        type="list",
        formula1="'Options'!$B$2:$B$6",
        allow_blank=False,
    )

    dv_misinfo = DataValidation(
        type="list",
        formula1="'Options'!$C$2:$C$5",
        allow_blank=False,
    )

    ws.add_data_validation(dv_frame)
    ws.add_data_validation(dv_stance)
    ws.add_data_validation(dv_misinfo)

    dv_frame.add("E5:E12")
    dv_stance.add("F5:F12")
    dv_misinfo.add("G5:G12")

    ws.freeze_panes = "A5"

    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 90
    ws.column_dimensions["E"].width = 38
    ws.column_dimensions["F"].width = 38
    ws.column_dimensions["G"].width = 38
    ws.column_dimensions["H"].width = 42

    for r in range(5, 13):
        ws.row_dimensions[r].height = 110

    wb.save(path)


build_workbook(
    OUT_A,
    "A",
)

build_workbook(
    OUT_B,
    "B",
)


print(
    "EXCERPT_MISMATCH_COUNT_A="
    + str(len(mismatch_a))
)

print(
    "EXCERPT_MISMATCH_COUNT_B="
    + str(len(mismatch_b))
)

print(
    "MISMATCH_IDS="
    + ",".join(mismatch_a)
)

print(
    "A_B_MISMATCH_SET_IDENTICAL=YES"
)

print(
    "TITLES_MATCH_FROZEN_PACKET=YES"
)

print(
    "PREVIOUS_LABELS_INCLUDED_IN_RECHECK=NO"
)

print(
    "RECHECK_ROWS_PER_ANNOTATOR="
    + str(len(recheck))
)

print(
    "EXACT_TEXT_RECHECK_BUILD=PASS"
)
